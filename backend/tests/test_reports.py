"""报表与导出用例：汇总口径、统计与 Excel 导出。"""

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

EQUIPMENT = "/api/v1/equipment"
OPERATIONS = "/api/v1/well-operations"
FAULTS = "/api/v1/fault-records"
REPORTS = "/api/v1/reports"


def _equipment(client: TestClient, headers: dict[str, str], no: str, status: str) -> dict:
    response = client.post(EQUIPMENT, headers=headers, json={"equipment_no": no, "status": status})
    assert response.status_code == 200, response.text
    return response.json()["data"]


def test_dashboard_matches_lists(api_client: TestClient, auth_headers: dict[str, str]) -> None:
    equipment = _equipment(api_client, auth_headers, "PLT-RPT-1", "在用")
    api_client.post(
        OPERATIONS,
        headers=auth_headers,
        json={
            "equipment_id": equipment["id"],
            "well_no": "井-RPT",
            "start_date": "2026-09-10 08:00:00",
            "end_date": "2026-09-10 16:00:00",
        },
    )
    api_client.post(
        FAULTS,
        headers=auth_headers,
        json={
            "equipment_id": equipment["id"],
            "occurred_at": "2026-09-11 10:00:00",
            "fault_type": "电气故障",
            "description": "报表用例",
        },
    )

    dashboard = api_client.get(f"{REPORTS}/dashboard", headers=auth_headers).json()["data"]
    equipment_list = api_client.get(EQUIPMENT, headers=auth_headers, params={"page_size": 100}).json()["data"]
    fault_list = api_client.get(FAULTS, headers=auth_headers, params={"page_size": 100}).json()["data"]

    assert dashboard["equipment_total"] == equipment_list["total"]
    assert sum(dashboard["status_counts"].values()) == equipment_list["total"]
    assert dashboard["open_faults"] == fault_list["total"] - sum(
        1 for item in fault_list["items"] if item["status"] == "已解决"
    )
    assert any(
        item["equipment_no"] == "PLT-RPT-1" and item["run_hours"] == "8.00"
        for item in dashboard["recent_operations"]
    )


def test_fault_statistics_agree_with_list(
    api_client: TestClient, auth_headers: dict[str, str]
) -> None:
    equipment = _equipment(api_client, auth_headers, "PLT-RPT-2", "在用")
    for index in range(2):
        api_client.post(
            FAULTS,
            headers=auth_headers,
            json={
                "equipment_id": equipment["id"],
                "occurred_at": f"2026-08-1{index + 1} 09:00:00",
                "fault_type": "密封失效",
                "description": f"统计用例 {index}",
            },
        )

    stats = api_client.get(
        f"{REPORTS}/fault-statistics",
        headers=auth_headers,
        params={"start": "2026-08-01", "end": "2026-08-31"},
    ).json()["data"]
    listed = api_client.get(
        FAULTS,
        headers=auth_headers,
        params={"fault_type": "密封失效", "page_size": 100},
    ).json()["data"]

    assert any(item["fault_type"] == "密封失效" and item["count"] >= 2 for item in stats["by_type"])
    assert stats["total"] >= 2
    assert stats["by_month"] and stats["by_month"][0]["month"].startswith("2026-0")
    assert listed["total"] >= 2


def test_equipment_utilization_reports_formula_and_hours(
    api_client: TestClient, auth_headers: dict[str, str]
) -> None:
    equipment = _equipment(api_client, auth_headers, "PLT-RPT-3", "在用")
    api_client.post(
        OPERATIONS,
        headers=auth_headers,
        json={
            "equipment_id": equipment["id"],
            "well_no": "井-USE",
            "start_date": "2026-09-15 06:00:00",
            "end_date": "2026-09-15 12:00:00",
        },
    )

    data = api_client.get(
        f"{REPORTS}/equipment-utilization",
        headers=auth_headers,
        params={"start": "2026-09-01", "end": "2026-09-30"},
    ).json()["data"]

    row = next(item for item in data["items"] if item["equipment_no"] == "PLT-RPT-3")
    assert row["operation_count"] == 1
    assert row["run_hours"] == "6.00"
    assert data["formula"]
    assert data["total_hours"] == "720.00"


def test_operation_trend_groups_by_month_and_instrument(
    api_client: TestClient, auth_headers: dict[str, str]
) -> None:
    equipment = _equipment(api_client, auth_headers, "PLT-RPT-4", "在用")
    api_client.post(
        OPERATIONS,
        headers=auth_headers,
        json={
            "equipment_id": equipment["id"],
            "start_date": "2026-09-18 08:00:00",
            "end_date": "2026-09-18 10:00:00",
            "instrument_type": "YST-48R",
        },
    )

    data = api_client.get(
        f"{REPORTS}/operation-trend", headers=auth_headers, params={"months": 24}
    ).json()["data"]

    assert data["by_month"]
    assert any(item["instrument_type"] == "YST-48R" for item in data["by_instrument"])


def test_export_operations_respects_filters(
    api_client: TestClient, auth_headers: dict[str, str]
) -> None:
    equipment = _equipment(api_client, auth_headers, "PLT-RPT-5", "在用")
    api_client.post(
        OPERATIONS,
        headers=auth_headers,
        json={
            "equipment_id": equipment["id"],
            "well_no": "井-导出",
            "start_date": "2026-09-19 08:00:00",
            "end_date": "2026-09-19 11:30:00",
            "operator_name": "张工",
        },
    )

    response = api_client.get(
        f"{OPERATIONS}/export", headers=auth_headers, params={"well_no": "井-导出"}
    )

    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert "attachment" in response.headers["content-disposition"]

    workbook = load_workbook(BytesIO(response.content), read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    assert rows[0][0] == "设备编号"
    assert rows[0][1] == "井号"
    assert len(rows) == 2  # 表头 + 一条记录
    assert rows[1][1] == "井-导出"  # 中文不乱码
    assert rows[1][9] == 3.5  # 使用时长自动算出 3.5 小时


def test_export_maintenance_orders(
    api_client: TestClient, auth_headers: dict[str, str]
) -> None:
    equipment = _equipment(api_client, auth_headers, "PLT-RPT-6", "维保")
    order = api_client.post(
        "/api/v1/maintenance-orders",
        headers=auth_headers,
        json={"equipment_id": equipment["id"], "level": "一级维保"},
    ).json()["data"]

    response = api_client.get(
        "/api/v1/maintenance-orders/export",
        headers=auth_headers,
        params={"equipment_id": equipment["id"]},
    )

    rows = list(load_workbook(BytesIO(response.content), read_only=True).active.iter_rows(values_only=True))
    assert rows[0][0] == "工单号"
    assert rows[1][0] == order["order_no"]
    assert rows[1][1] == "PLT-RPT-6"


def test_readonly_user_cannot_export(
    api_client: TestClient, readonly_headers: dict[str, str]
) -> None:
    response = api_client.get(f"{OPERATIONS}/export", headers=readonly_headers)

    assert response.status_code == 403
    assert response.json()["code"] == 40300
