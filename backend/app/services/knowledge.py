"""知识库业务：解析文档、切块入库、列表与删除。

解析支持四类格式：txt / md（直接读文本）、docx（python-docx）、pdf（pdfplumber）。
**扫描件 PDF 提取不到文字**，这时不静默入库，而是把文档标记为解析失败并写明原因——
知识库里混进空文档比报错更糟：用户以为传上去了，检索却永远不命中。

分块策略：先按行聚合到接近阈值，超长行按句末标点切；块与块之间留少量重叠，
避免一句话被切断后两边都检索不到。
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from uuid import uuid4

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import Settings, get_settings
from app.core.errors import BusinessError, ErrorCode
from app.core.timeutil import local_now
from app.models.knowledge import (
    SOURCE_FILE,
    SOURCE_TEXT,
    STATUS_FAILED,
    STATUS_READY,
    KnowledgeChunk,
    KnowledgeDocument,
)
from app.models.user import User

logger = logging.getLogger(__name__)

# 分块参数：500 字是中文技术文档的常见粒度，50 字重叠保证跨块句子可被检索到
CHUNK_SIZE = 500
CHUNK_OVERLAP = 50

SENTENCE_END = re.compile(r"(?<=[。！？；.!?;])")


# ----------------------------------------------------------------- 解析


def _decode_text(content: bytes) -> str:
    """文本文件解码：优先 UTF-8，退回 GBK（现场导出的 txt 常是 GBK）。"""
    for encoding in ("utf-8", "utf-8-sig", "gbk"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise BusinessError(ErrorCode.BUSINESS_RULE, "文本编码无法识别，请转成 UTF-8 后重试", 422)


def _read_docx(content: bytes) -> str:
    import io

    from docx import Document as DocxDocument

    document = DocxDocument(io.BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_pdf(content: bytes) -> str:
    import io

    import pdfplumber

    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)
    return "\n".join(parts)


def _check_extension(file_name: str, settings: Settings) -> str:
    """扩展名在落盘之前就校验：不认识的格式不该在服务器上留下原件。"""
    suffix = Path(file_name).suffix.lower()
    allowed = {item.lower() for item in settings.knowledge_upload_extensions}
    if suffix not in allowed:
        raise BusinessError(
            ErrorCode.BUSINESS_RULE,
            f"知识库不支持 {suffix or '未知'} 格式，可上传 {'、'.join(sorted(allowed))}",
            422,
        )
    return suffix


def _check_size(content: bytes, settings: Settings) -> None:
    limit = settings.upload_max_size_mb * 1024 * 1024
    if len(content) > limit:
        raise BusinessError(
            ErrorCode.BUSINESS_RULE,
            f"文件大小 {len(content) / 1024 / 1024:.1f} MB 超过限制（最大 {settings.upload_max_size_mb} MB）",
            413,
        )


def extract_text(file_name: str, content: bytes, settings: Settings | None = None) -> str:
    """按扩展名抽取正文；不支持的格式直接报错并把可选项告诉用户。"""
    settings = settings or get_settings()
    suffix = _check_extension(file_name, settings)

    if suffix in {".txt", ".md", ".markdown"}:
        text = _decode_text(content)
    elif suffix == ".docx":
        text = _read_docx(content)
    elif suffix == ".pdf":
        text = _read_pdf(content)
    else:  # pragma: no cover - allowed 已过滤，这里是兜底
        raise BusinessError(ErrorCode.BUSINESS_RULE, f"暂不支持解析 {suffix}", 422)

    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        raise BusinessError(
            ErrorCode.BUSINESS_RULE,
            "没有提取到任何文字：如果是扫描件 PDF，请提供文字版 PDF 或 Word 文档",
            422,
        )
    return text


# ----------------------------------------------------------------- 分块


def _split_long_line(line: str, size: int) -> list[str]:
    """超长行按句末标点切；仍超长就硬切。"""
    pieces: list[str] = []
    current = ""
    for sentence in SENTENCE_END.split(line):
        if not sentence:
            continue
        if len(current) + len(sentence) <= size:
            current += sentence
            continue
        if current:
            pieces.append(current)
        while len(sentence) > size:
            pieces.append(sentence[:size])
            sentence = sentence[size:]
        current = sentence
    if current:
        pieces.append(current)
    return pieces


def split_chunks(text: str, *, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """把正文切成若干块，块间保留 overlap 个字符的重叠。"""
    units: list[str] = []
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        units.extend(_split_long_line(line, size) if len(line) > size else [line])

    chunks: list[str] = []
    current = ""
    for unit in units:
        if current and len(current) + len(unit) + 1 > size:
            chunks.append(current)
            tail = current[-overlap:] if overlap > 0 else ""
            current = f"{tail}{unit}" if tail else unit
        else:
            current = f"{current}\n{unit}" if current else unit
    if current:
        chunks.append(current)
    return chunks


# ----------------------------------------------------------------- 入库


def _store_file(content: bytes, suffix: str, settings: Settings) -> str:
    settings.knowledge_path.mkdir(parents=True, exist_ok=True)
    stored_name = f"{uuid4().hex}{suffix}"
    (settings.knowledge_path / stored_name).write_bytes(content)
    return stored_name


def _persist(
    db: Session,
    *,
    title: str,
    source_type: str,
    user: User,
    file_name: str | None = None,
    stored_name: str | None = None,
    content_type: str | None = None,
    size_bytes: int | None = None,
) -> KnowledgeDocument:
    document = KnowledgeDocument(
        title=title[:255],
        source_type=source_type,
        file_name=file_name,
        stored_name=stored_name,
        content_type=content_type,
        size_bytes=size_bytes,
        status=STATUS_FAILED,  # 先按失败落库，解析成功后再改 ready
        char_count=0,
        chunk_count=0,
        is_deleted=False,
        created_by=user.id,
        updated_by=user.id,
    )
    db.add(document)
    db.flush()
    return document


def create_from_text(db: Session, *, title: str, content: str, user: User) -> KnowledgeDocument:
    """页面粘贴的文本：直接分块入库。"""
    text = (content or "").strip()
    if not text:
        raise BusinessError(ErrorCode.BUSINESS_RULE, "正文不能为空", 422)

    document = _persist(db, title=title, source_type=SOURCE_TEXT, user=user)
    _fill_chunks(db, document, text)
    return document


def create_from_file(
    db: Session,
    *,
    file_name: str,
    content: bytes,
    content_type: str | None,
    user: User,
    title: str | None = None,
    settings: Settings | None = None,
) -> KnowledgeDocument:
    """上传文件：先落盘保留原件，再解析分块；解析失败也留一条记录（页面能看到原因）。"""
    settings = settings or get_settings()
    # 扩展名与大小先校验，不合格的请求不落盘、不入库，直接 4xx 打回
    suffix = _check_extension(file_name, settings)
    _check_size(content, settings)

    stored_name = _store_file(content, suffix, settings)
    document = _persist(
        db,
        title=title or Path(file_name).stem,
        source_type=SOURCE_FILE,
        user=user,
        file_name=file_name,
        stored_name=stored_name,
        content_type=content_type,
        size_bytes=len(content),
    )

    try:
        text = extract_text(file_name, content, settings)
    except BusinessError as exc:
        document.error = exc.message[:255]
        logger.warning("知识库文档解析失败：%s（%s）", file_name, exc.message)
        return document

    _fill_chunks(db, document, text)
    return document


def _fill_chunks(db: Session, document: KnowledgeDocument, text: str) -> None:
    chunks = split_chunks(text)
    for seq, chunk in enumerate(chunks, start=1):
        db.add(
            KnowledgeChunk(
                document_id=document.id,
                seq=seq,
                text=chunk,
                char_count=len(chunk),
            )
        )
    document.char_count = len(text)
    document.chunk_count = len(chunks)
    document.status = STATUS_READY
    document.error = None
    document.updated_by = document.created_by
    db.flush()


# ----------------------------------------------------------------- 查询与删除


def get_document(db: Session, document_id: int, *, include_deleted: bool = False) -> KnowledgeDocument:
    document = db.get(KnowledgeDocument, document_id)
    if document is None or (document.is_deleted and not include_deleted):
        raise BusinessError(ErrorCode.NOT_FOUND, "文档不存在或已删除", 404)
    return document


def list_documents(
    db: Session,
    *,
    keyword: str | None = None,
    status: str | None = None,
    page: int = 1,
    page_size: int = 20,
) -> tuple[list[KnowledgeDocument], int]:
    stmt = select(KnowledgeDocument).where(KnowledgeDocument.is_deleted.is_(False))
    if keyword:
        stmt = stmt.where(KnowledgeDocument.title.like(f"%{keyword.strip()}%"))
    if status:
        stmt = stmt.where(KnowledgeDocument.status == status)

    total = db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    items = list(
        db.execute(
            stmt.order_by(KnowledgeDocument.id.desc())
            .limit(page_size)
            .offset((page - 1) * page_size)
        ).scalars()
    )
    return items, total


def list_chunks(db: Session, document_id: int, *, limit: int = 3) -> list[KnowledgeChunk]:
    """取前几个分块，用于详情预览。"""
    return list(
        db.execute(
            select(KnowledgeChunk)
            .where(KnowledgeChunk.document_id == document_id)
            .order_by(KnowledgeChunk.seq)
            .limit(limit)
        ).scalars()
    )


def delete_document(db: Session, document_id: int, user: User) -> KnowledgeDocument:
    """软删除：内容可恢复，审计能查到删除人。"""
    document = get_document(db, document_id)
    document.is_deleted = True
    document.deleted_at = local_now()
    document.updated_by = user.id
    db.flush()
    return document


def stats(db: Session) -> dict:
    """知识库概览：文档数、分块数、解析失败数。"""
    active = KnowledgeDocument.is_deleted.is_(False)
    documents = db.scalar(select(func.count()).select_from(KnowledgeDocument).where(active)) or 0
    failed = (
        db.scalar(
            select(func.count())
            .select_from(KnowledgeDocument)
            .where(active, KnowledgeDocument.status == STATUS_FAILED)
        )
        or 0
    )
    chunks = (
        db.scalar(
            select(func.count())
            .select_from(KnowledgeChunk)
            .join(KnowledgeDocument, KnowledgeChunk.document_id == KnowledgeDocument.id)
            .where(active)
        )
        or 0
    )
    characters = (
        db.scalar(
            select(func.coalesce(func.sum(KnowledgeDocument.char_count), 0)).where(active)
        )
        or 0
    )
    return {
        "documents": documents,
        "chunks": chunks,
        "failed": failed,
        "characters": int(characters),
    }
