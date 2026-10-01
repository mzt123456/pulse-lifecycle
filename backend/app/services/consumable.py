"""耗材目录、批量导入与领用扣减。

领用扣减先对耗材行加排他锁，再校验库存，避免并发领用把库存扣成负数
（编写报告 4.5）。
"""

from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.errors import BusinessError, ErrorCode
from app.core.timeutil import local_now
from app.models.consumable import Consumable, ConsumableUsage
from app.models.maintenance import MaintenanceOrder
from app.models.user import User
from app.schemas.consumable import (
    ConsumableCreate,
    ConsumableUpdate,
    ConsumableUsageCreate,
    ImportResult,
)

IMPORT_HEADERS = ["部件号", "中文描述", "库存数量", "安全库存", "单位"]


def list_consumables(
    db: Session,
    keyword: str | None = None,
    low_stock_only: bool = False,
    page: int = 1,
    page_size: int = 20,
) -> tuple[list[Consumable], int]:
    stmt = select(Consumable)
    if keyword:
        pattern = f"%{keyword.strip()}%"
        stmt = stmt.where(
            Consumable.part_no.like(pattern) | Consumable.description.like(pattern)
        )
    if low_stock_only:
        stmt = stmt.where(Consumable.quantity <= Consumable.safety_stock)

    total = db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    items = list(
        db.execute(
            stmt.order_by(Consumable.part_no).limit(page_size).offset((page - 1) * page_size)
        ).scalars()
    )
    return items, total


def get_consumable(db: Session, consumable_id: int) -> Consumable:
    item = db.get(Consumable, consumable_id)
    if item is None:
        raise BusinessError(ErrorCode.NOT_FOUND, "耗材不存在", 404)
    return item


def create_consumable(db: Session, payload: ConsumableCreate, user: User) -> Consumable:
    existing = db.scalar(select(Consumable.id).where(Consumable.part_no == payload.part_no))
    if existing is not None:
        raise BusinessError(ErrorCode.CONFLICT, f"部件号 {payload.part_no} 已存在", 409)

    item = Consumable(**payload.model_dump(), created_by=user.id, updated_by=user.id)
    db.add(item)
    db.flush()
    return item


def update_consumable(
    db: Session, consumable_id: int, payload: ConsumableUpdate, user: User
) -> Consumable:
    item = get_consumable(db, consumable_id)
    for field, value in payload.model_dump(exclude_unset=True).items():
        setattr(item, field, value)
    item.updated_by = user.id
    db.flush()
    return item


def issue_consumable(
    db: Session, payload: ConsumableUsageCreate, user: User
) -> ConsumableUsage:
    """领用耗材：加行级锁后校验库存，库存不足直接拒绝。"""
    item = db.execute(
        select(Consumable).where(Consumable.id == payload.consumable_id).with_for_update()
    ).scalar_one_or_none()
    if item is None:
        raise BusinessError(ErrorCode.NOT_FOUND, "耗材不存在", 404)

    if payload.order_id is not None and db.get(MaintenanceOrder, payload.order_id) is None:
        raise BusinessError(ErrorCode.NOT_FOUND, "关联维保工单不存在", 404)

    if item.quantity < payload.quantity:
        raise BusinessError(
            ErrorCode.BUSINESS_RULE,
            f"库存不足：{item.description} 当前 {item.quantity}，本次领用 {payload.quantity}",
            422,
        )

    item.quantity -= payload.quantity
    item.updated_by = user.id

    usage = ConsumableUsage(
        consumable_id=item.id,
        order_id=payload.order_id,
        quantity=payload.quantity,
        used_by=user.id,
        used_at=local_now(),
        remark=payload.remark,
    )
    db.add(usage)
    db.flush()
    db.expire(usage, ["consumable", "used_by_user"])
    return usage


def list_usages(
    db: Session,
    consumable_id: int | None = None,
    page: int = 1,
    page_size: int = 20,
) -> tuple[list[ConsumableUsage], int]:
    stmt = select(ConsumableUsage)
    if consumable_id is not None:
        stmt = stmt.where(ConsumableUsage.consumable_id == consumable_id)

    total = db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    items = list(
        db.execute(
            stmt.order_by(ConsumableUsage.used_at.desc(), ConsumableUsage.id.desc())
            .limit(page_size)
            .offset((page - 1) * page_size)
        )
        .unique()
        .scalars()
    )
    return items, total


def import_consumables(db: Session, content: bytes, user: User) -> ImportResult:
    """按 Excel 批量导入目录，以部件号为唯一键：存在则更新，不存在则新增。"""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # 文件不是有效的 xlsx
        raise BusinessError(ErrorCode.BUSINESS_RULE, f"无法解析 Excel 文件：{exc}", 422) from exc

    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    except StopIteration:
        raise BusinessError(ErrorCode.BUSINESS_RULE, "Excel 文件为空", 422) from None

    missing = [name for name in IMPORT_HEADERS if name not in header]
    if missing:
        raise BusinessError(
            ErrorCode.BUSINESS_RULE,
            f"缺少必需列：{'、'.join(missing)}（应包含 {'、'.join(IMPORT_HEADERS)}）",
            422,
        )
    index = {name: header.index(name) for name in IMPORT_HEADERS}

    result = ImportResult()
    for line_no, row in enumerate(rows, start=2):
        if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        part_no = str(row[index["部件号"]]).strip() if row[index["部件号"]] is not None else ""
        description = (
            str(row[index["中文描述"]]).strip() if row[index["中文描述"]] is not None else ""
        )
        if not part_no or not description:
            result.failed += 1
            result.errors.append(f"第 {line_no} 行：部件号与中文描述必填")
            continue

        def to_int(value: object) -> int:
            if value is None or str(value).strip() == "":
                return 0
            return int(float(str(value)))

        try:
            quantity = to_int(row[index["库存数量"]])
            safety_stock = to_int(row[index["安全库存"]])
        except ValueError:
            result.failed += 1
            result.errors.append(f"第 {line_no} 行：库存数量或安全库存不是数字")
            continue

        unit = row[index["单位"]]
        item = db.scalar(select(Consumable).where(Consumable.part_no == part_no))
        if item is None:
            db.add(
                Consumable(
                    part_no=part_no,
                    description=description,
                    quantity=quantity,
                    safety_stock=safety_stock,
                    unit=str(unit).strip() if unit is not None else None,
                    created_by=user.id,
                    updated_by=user.id,
                )
            )
            result.created += 1
        else:
            item.description = description
            item.quantity = quantity
            item.safety_stock = safety_stock
            item.unit = str(unit).strip() if unit is not None else None
            item.updated_by = user.id
            result.updated += 1

    db.flush()
    return result
